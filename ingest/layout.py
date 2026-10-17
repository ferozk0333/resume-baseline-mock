"""Layout reader for code-first extraction (specs/18 §4.2 item 14): PDF/DOCX → visual lines in reading order.

Each line keeps its position, font size, bold, colour and column, so parsers can tell a heading from a bullet
and a sidebar from the main column. Hidden text (white or near-background colour, tiny font, off-page,
invisible) is found here and never reaches a parser; the same check removes it from the source text when
REMOVE_HIDDEN_TEXT is on (spec item 20). The pymupdf4llm Markdown stays the source text shown in Review and used
by the quote check; this reader only decides structure.
"""
from __future__ import annotations

import html
import io
import re
from dataclasses import dataclass, field

import pymupdf

from core import config

BULLET_GLYPHS = "•▪◦●‣■◆◇➢➤►▸✓✔○□–-*·"
_BULLET = re.compile(rf"^\s*([{re.escape(BULLET_GLYPHS)}]|[-])\s+")
_GAP = re.compile(r"(?:\t+| {3,})")
_FOOTER = re.compile(r"^Page \d+ of \d+$")


@dataclass
class Hidden:
    """Text removed as hidden: kept for Admin, never shown to a parser or a model."""
    page: int
    text: str
    reason: str                     # colour | tiny | off_page | invisible | docx_hidden
    bbox: tuple[float, float, float, float] | None = None


@dataclass
class Line:
    text: str                       # cells joined by one space, whitespace collapsed
    cells: list[str]                # parts split at tab stops / wide gaps (e.g. a right-aligned date)
    page: int
    x0: float
    y0: float
    x1: float
    y1: float
    size: float
    bold: bool
    color: int                      # 0xRRGGBB
    block: int                      # PDF text block (-1 for DOCX)
    column: int = 0                 # 0 = full width or single column; 1 = left; 2 = right
    bullet: bool = False
    style: str = ""                 # DOCX paragraph style name
    cell: int | None = None         # DOCX table cell index
    lead_bold: bool = False         # starts with a bold run ("Languages: …" in a skills section)
    bold_lead: str = ""             # the leading bold text itself (a project's name before its description)


@dataclass
class Layout:
    kind: str                                   # pdf | docx
    lines: list[Line] = field(default_factory=list)
    hidden: list[Hidden] = field(default_factory=list)
    pages: int = 1
    columns: int = 1                            # most columns seen on one page
    page_width: float = 612.0
    footers: int = 0                            # "Page N of M" lines (LinkedIn's export has one per page)
    tables: int = 0                             # rows read across table cells (dates in one cell, the entry in another)

    @property
    def body_size(self) -> float:
        """The most common font size (weighted by characters): body text."""
        weights: dict[float, int] = {}
        for ln in self.lines:
            weights[round(ln.size, 1)] = weights.get(round(ln.size, 1), 0) + len(ln.text)
        return max(weights, key=weights.get) if weights else 10.0

    def text(self) -> str:
        return "\n".join(("- " if ln.bullet else "") + ln.text for ln in self.lines)


# --- Hidden text -------------------------------------------------------------------------------------------------

def _rgb(value) -> tuple[float, float, float]:
    if isinstance(value, int):
        return ((value >> 16) & 255) / 255, ((value >> 8) & 255) / 255, (value & 255) / 255
    if value is None:
        return 1.0, 1.0, 1.0
    if len(value) == 1:                         # grey
        return value[0], value[0], value[0]
    if len(value) == 4:                         # CMYK
        c, m, y, k = value
        return (1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)
    return tuple(value[:3])


def _luminance(rgb) -> float:
    def ch(c):
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (ch(c) for c in rgb)
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def contrast(a, b) -> float:
    """WCAG contrast ratio of two colours (1 = identical, 21 = black on white)."""
    la, lb = sorted((_luminance(_rgb(a)), _luminance(_rgb(b))), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _fills(page) -> list[tuple[pymupdf.Rect, tuple]]:
    """Filled shapes in paint order (later ones are on top): what text could sit on."""
    out = []
    for d in page.get_drawings():
        if d.get("fill") is not None and "f" in (d.get("type") or "") and (d.get("fill_opacity") or 1) >= 0.5:
            out.append((pymupdf.Rect(d["rect"]), d["fill"]))
    return out


def _background(fills, bbox) -> tuple:
    r = pymupdf.Rect(bbox)
    center = pymupdf.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
    for rect, fill in reversed(fills):
        if center in rect:
            return fill
    return (1.0, 1.0, 1.0)


def hidden_reason(span: dict, page_rect, fills) -> str | None:
    """Why a PDF span cannot be seen, or None when it is visible."""
    r = pymupdf.Rect(span["bbox"])
    if not r.intersects(page_rect) or not page_rect.contains(pymupdf.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)):
        return "off_page"
    if span.get("alpha", 255) == 0:
        return "invisible"
    if span["size"] < config.HIDDEN_MIN_FONT_PT:
        return "tiny"
    if contrast(span["color"], _background(fills, span["bbox"])) < config.HIDDEN_MIN_CONTRAST:
        return "colour"
    return None


def pdf_hidden_spans(page, pno: int) -> list[Hidden]:
    fills = _fills(page)
    out = []
    for b in page.get_text("dict")["blocks"]:
        for ln in b.get("lines", []):
            for s in ln["spans"]:
                if s["text"].strip() and (why := hidden_reason(s, page.rect, fills)):
                    out.append(Hidden(pno, s["text"].strip(), why, tuple(s["bbox"])))
    return out


def remove_pdf_hidden(doc) -> list[Hidden]:
    """Redact hidden spans in place (text only; images and drawings stay). Returns what was removed."""
    removed: list[Hidden] = []
    for pno, page in enumerate(doc):
        spans = pdf_hidden_spans(page, pno)
        for h in spans:
            r = pymupdf.Rect(h.bbox)
            if r.is_empty or not r.intersects(page.rect):
                r = r & page.rect                       # off-page text: redact the part inside, if any
            if r.is_empty:
                continue
            mid = (r.y0 + r.y1) / 2                     # a thin band: never touch visible text on the next line
            page.add_redact_annot(pymupdf.Rect(r.x0, mid - max(0.5, r.height / 4), r.x1, mid + max(0.5, r.height / 4)))
        if spans:
            page.apply_redactions(images=0, graphics=0, text=0)   # PDF_REDACT_IMAGE_NONE, LINE_ART_NONE, TEXT_REMOVE
        removed += spans
    return removed


# --- PDF ---------------------------------------------------------------------------------------------------------

def _clean(text: str) -> str:
    text = html.unescape(text).replace(" ", " ").replace("�", " ")
    return re.sub(r"\s+", " ", text).strip()


def _pdf_line(ln: dict, pno: int, block: int, page_rect, fills, hidden: list[Hidden]) -> Line | None:
    spans = []
    for s in ln["spans"]:
        if not s["text"].strip():
            spans.append(s)
            continue
        why = hidden_reason(s, page_rect, fills)
        if why:
            hidden.append(Hidden(pno, s["text"].strip(), why, tuple(s["bbox"])))
        else:
            spans.append(s)
    visible = [s for s in spans if s["text"].strip()]
    if not visible:
        return None
    # Cells: a tab-stop run of spaces inside a span, or a wide horizontal jump between spans.
    cells, cur, last_x1 = [], "", None
    for s in spans:
        jump = last_x1 is not None and s["bbox"][0] - last_x1 > 3 * max(s["size"], 1) and s["text"].strip()
        if jump and cur.strip():
            cells.append(cur)
            cur = ""
        parts = _GAP.split(s["text"])
        cur += parts[0]
        for p in parts[1:]:
            if cur.strip():
                cells.append(cur)
            cur = p
        if s["text"].strip():
            last_x1 = s["bbox"][2]
    if cur.strip():
        cells.append(cur)
    cells = [c for c in (_clean(c) for c in cells) if c]
    chars = sum(len(s["text"].strip()) for s in visible)
    bold_chars = sum(len(s["text"].strip()) for s in visible if s["flags"] & 16 or "bold" in s["font"].lower())
    main = max(visible, key=lambda s: len(s["text"].strip()))
    x0 = min(s["bbox"][0] for s in visible)
    first = next(s for s in visible if s["text"].strip() and not _BULLET.match(s["text"] + " ")) \
        if any(not _BULLET.match(s["text"] + " ") for s in visible) else visible[0]
    line = Line(text=" ".join(cells), cells=cells, page=pno, x0=x0, y0=min(s["bbox"][1] for s in visible),
                x1=max(s["bbox"][2] for s in visible), y1=max(s["bbox"][3] for s in visible), size=main["size"],
                bold=bold_chars >= 0.6 * chars, color=main["color"], block=block,
                lead_bold=bool(first["flags"] & 16 or "bold" in first["font"].lower()))
    lead = ""
    for s in visible[visible.index(first):]:
        if not (s["flags"] & 16 or "bold" in s["font"].lower()):
            break
        lead += s["text"]
    line.bold_lead = _clean(lead) if line.lead_bold else ""
    _mark_bullet(line)
    return line if line.text else None


def _mark_bullet(line: Line) -> None:
    m = _BULLET.match(line.text)
    if m and len(line.text) > len(m.group(0)):
        line.bullet = True
        line.text = line.text[m.end():].strip()
        if line.cells:
            line.cells[0] = _BULLET.sub("", line.cells[0], count=1).strip() or line.cells[0]
            line.cells = [c for c in line.cells if c]


def _columns(lines: list[Line], width: float) -> int:
    """Find one vertical gutter (a gap no line crosses, with real text on both sides) and mark columns."""
    if len(lines) < 6:
        return 1
    xs = sorted({round(x) for ln in lines for x in (ln.x0, ln.x1)})
    best = None
    lo, hi = width * 0.15, width * 0.85
    for x in range(int(lo), int(hi), 2):
        crossing = [ln for ln in lines if ln.x0 < x - 1 < ln.x1 or ln.x0 < x + 1 < ln.x1]
        left = [ln for ln in lines if ln.x1 <= x]
        right = [ln for ln in lines if ln.x0 >= x]
        if len(left) < 3 or len(right) < 3 or len(crossing) > max(3, len(lines) // 8):
            continue
        # both sides must run alongside each other, not one above the other
        overlap = min(max(ln.y1 for ln in left), max(ln.y1 for ln in right)) - \
            max(min(ln.y0 for ln in left), min(ln.y0 for ln in right))
        if overlap <= 0:
            continue
        score = min(len(left), len(right)) - 3 * len(crossing)
        if best is None or score > best[0]:
            best = (score, x)
    del xs
    if best is None:
        return 1
    gutter = best[1]
    # A table (dates in a narrow left cell, the entry on the right) is rows, not two columns: most left lines sit
    # on the same row as a right line. Read it row by row.
    left = [ln for ln in lines if ln.x1 <= gutter + 1]
    right = [ln for ln in lines if ln.x0 >= gutter - 1]
    # A table's left cells that share a row with the right side are dates; a sidebar's are ordinary text.
    paired = [a for a in left if any(abs(a.y0 - b.y0) < 0.6 * a.size for b in right)]
    dated = [a for a in paired if re.search(r"\b(19|20)\d{2}\b", a.text) and len(a.text.split()) <= 6]
    if len(paired) >= 2 and len(dated) >= 0.6 * len(paired):
        return 1
    for ln in lines:
        ln.column = 1 if ln.x1 <= gutter + 1 else 2 if ln.x0 >= gutter - 1 else 0
    return 2


def _rows(lines: list[Line]) -> list[Line]:
    """One column: pieces of one visual row in different blocks (a table cell, a right-aligned date) become one line
    with cells, left to right."""
    out: list[Line] = []
    first_cell: dict[int, float] = {}                 # id(row line) → right edge of its first cell
    for ln in lines:
        prev = out[-1] if out else None
        if prev is not None and prev.page == ln.page and ln.x0 >= prev.x1 - 1 and \
                abs((prev.y0 + prev.y1) / 2 - (ln.y0 + ln.y1) / 2) < 0.4 * max(ln.size, prev.size) and not ln.bullet:
            first_cell.setdefault(id(prev), prev.x1)
            if not prev.bold_lead and ln.bold_lead:
                prev.bold_lead, prev.lead_bold = ln.bold_lead, True
            prev.cells += ln.cells
            prev.text = " ".join(prev.cells)
            prev.x1, prev.y1 = ln.x1, max(prev.y1, ln.y1)
            continue
        edge = first_cell.get(id(prev)) if prev is not None else None
        if edge is not None and ln.x1 <= edge + 2 and not ln.bullet and len(ln.text.split()) <= 3 and \
                ln.y0 - prev.y0 < 2.5 * ln.size:
            prev.cells[0] = f"{prev.cells[0]} {ln.text}"   # the first cell wrapped ("Jul 2024 – / Present")
            prev.text = " ".join(prev.cells)
            continue
        out.append(ln)
    return out


def _reading_order(lines: list[Line]) -> list[Line]:
    cols = [ln for ln in lines if ln.column]
    if not cols:
        return sorted(lines, key=lambda ln: (round(ln.y0, 1), ln.x0))
    # Both columns start at `top`; anything above it (name, contact line) is the header, read first.
    top = max(min((ln.y0 for ln in cols if ln.column == c), default=0) for c in (1, 2))
    head = sorted((ln for ln in lines if ln.y0 < top - 1), key=lambda ln: (ln.y0, ln.x0))
    rest = sorted((ln for ln in lines if not ln.column and ln.y0 >= top - 1), key=lambda ln: (ln.y0, ln.x0))
    left = sorted((ln for ln in cols if ln.column == 1 and ln.y0 >= top - 1), key=lambda ln: (ln.y0, ln.x0))
    right = sorted((ln for ln in cols if ln.column == 2 and ln.y0 >= top - 1), key=lambda ln: (ln.y0, ln.x0))
    return head + left + right + rest


def read_pdf(data: bytes) -> Layout:
    doc = pymupdf.open(stream=data, filetype="pdf")
    out = Layout("pdf", pages=len(doc), page_width=doc[0].rect.width if len(doc) else 612.0)
    for pno, page in enumerate(doc):
        fills = _fills(page)
        lines: list[Line] = []
        for b in page.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            for ln in b["lines"]:
                line = _pdf_line(ln, pno, b["number"], page.rect, fills, out.hidden)
                if line is None:
                    continue
                if _FOOTER.match(line.text):
                    out.footers += 1
                    continue
                prev = lines[-1] if lines else None
                if prev is not None and prev.block == line.block and line.x0 >= prev.x1 and \
                        abs((prev.y0 + prev.y1) / 2 - (line.y0 + line.y1) / 2) < 0.4 * line.size:
                    prev.cells += line.cells                # same baseline: a right-aligned date or location
                    prev.text = " ".join(prev.cells)
                    prev.x1, prev.y1 = line.x1, max(prev.y1, line.y1)
                    continue
                lines.append(line)
        cols = _columns(lines, page.rect.width)
        out.columns = max(out.columns, cols)
        ordered = _reading_order(lines)
        if cols == 1:
            ordered = _rows(ordered)
        out.lines += ordered
    return out


# --- DOCX --------------------------------------------------------------------------------------------------------

def _run_hidden(run) -> str | None:
    f = run.font
    if f.hidden:
        return "docx_hidden"
    if f.size is not None and f.size.pt < config.HIDDEN_MIN_FONT_PT:
        return "tiny"
    try:
        rgb = f.color.rgb if f.color is not None and f.color.type is not None else None
    except (AttributeError, ValueError):
        rgb = None
    if rgb is not None and contrast(int(str(rgb), 16), (1.0, 1.0, 1.0)) < config.HIDDEN_MIN_CONTRAST:
        return "colour"
    return None


def docx_runs(paragraph):
    """Runs in order, including those inside hyperlinks."""
    from docx.text.hyperlink import Hyperlink
    for item in paragraph.iter_inner_content():
        if isinstance(item, Hyperlink):
            yield from item.runs
        else:
            yield item


def docx_visible_text(paragraph, hidden: list[Hidden] | None = None) -> str:
    parts = []
    for run in docx_runs(paragraph):
        why = _run_hidden(run)
        if why:
            if hidden is not None and run.text.strip():
                hidden.append(Hidden(0, run.text.strip(), why))
            continue
        parts.append(run.text)
    return "".join(parts)


def _docx_line(p, hidden: list[Hidden], cell: int | None = None) -> Line | None:
    raw = docx_visible_text(p, hidden)
    if not raw.strip():
        return None
    style = (p.style.name if p.style is not None else "") or ""
    runs = [r for r in docx_runs(p) if r.text.strip() and not _run_hidden(r)]
    chars = sum(len(r.text.strip()) for r in runs) or 1
    style_bold = bool(p.style is not None and p.style.font is not None and p.style.font.bold)
    bold = sum(len(r.text.strip()) for r in runs if r.bold or (r.bold is None and style_bold)) >= 0.6 * chars
    sizes = [r.font.size.pt for r in runs if r.font.size is not None]
    size = max(sizes) if sizes else (p.style.font.size.pt if p.style is not None and p.style.font is not None
                                     and p.style.font.size is not None else 11.0)
    heading = re.match(r"Heading (\d)", style)
    if style == "Title":
        size, bold = max(size, 26.0), True
    elif heading:
        size, bold = max(size, 18.0 - 2 * int(heading.group(1))), True
    cells = [c for c in (_clean(c) for c in _GAP.split(raw)) if c]
    line = Line(text=" ".join(cells), cells=cells, page=0, x0=0, y0=0, x1=0, y1=0, size=size, bold=bold, color=0,
                block=-1, style=style, cell=cell)
    lead = ""
    for r in runs:
        if not (r.bold or (r.bold is None and style_bold)):
            break
        lead += r.text
    line.lead_bold, line.bold_lead = bool(lead.strip()), _clean(lead)
    numbered = p._p.pPr is not None and p._p.pPr.numPr is not None
    if style.startswith("List") or numbered:
        line.bullet = True
    _mark_bullet(line)
    return line if line.text else None


def read_docx(data: bytes) -> Layout:
    import docx
    document = docx.Document(io.BytesIO(data))
    out = Layout("docx")
    for block in document.iter_inner_content():
        if isinstance(block, docx.table.Table):
            out.tables += len(block.rows)
            for row in block.rows:
                seen = set()
                for i, c in enumerate(row.cells):
                    if id(c._tc) in seen:                 # merged cells repeat
                        continue
                    seen.add(id(c._tc))
                    for p in c.paragraphs:
                        if (ln := _docx_line(p, out.hidden, cell=i)) is not None:
                            out.lines.append(ln)
        elif (ln := _docx_line(block, out.hidden)) is not None:
            out.lines.append(ln)
    return out


def read_layout(filename: str, data: bytes) -> Layout:
    if filename.lower().endswith(".pdf"):
        return read_pdf(data)
    if filename.lower().endswith(".docx"):
        return read_docx(data)
    raise ValueError("Upload a PDF or DOCX file.")


# --- Logical lines: wrapped lines and hyphenation joined -------------------------------------------------------------

def _column_right(lines: list[Line]) -> dict[tuple[int, int], tuple[float, float]]:
    out: dict[tuple[int, int], tuple[float, float]] = {}
    for ln in lines:
        k = (ln.page, ln.column)
        lo, hi = out.get(k, (ln.x0, ln.x1))
        out[k] = (min(lo, ln.x0), max(hi, ln.x1))
    return out


def _join(a: str, b: str, source: str | None) -> str:
    """Join a wrapped line. A word broken by a hyphen is joined the way the source text has it."""
    if re.search(r"[A-Za-z]-$", a) and re.match(r"[a-z]", b):
        word_a, word_b = re.search(r"(\w+)-$", a).group(1), re.match(r"\w+", b).group(0)
        joined = f"{word_a}{word_b}"
        if source is not None and joined.casefold() in source.casefold() and \
                f"{word_a}-{word_b}".casefold() not in source.casefold():
            return a[:-1] + b
        return a + b
    return f"{a} {b}"


def logical_lines(layout: Layout, source: str | None = None) -> list[Line]:
    """Merge wrapped continuation lines into the line they continue (bullets and paragraphs)."""
    if layout.kind == "docx":
        return [Line(**{**ln.__dict__, "cells": list(ln.cells)}) for ln in layout.lines]
    bounds = _column_right(layout.lines)
    out: list[Line] = []
    for ln in layout.lines:
        prev = out[-1] if out else None
        if prev is not None and _continues(prev, ln, bounds):
            prev.text = _join(prev.text, ln.text, source)
            prev.cells = [prev.text]
            prev.x1, prev.y1 = max(prev.x1, ln.x1), ln.y1
            continue
        out.append(Line(**{**ln.__dict__, "cells": list(ln.cells)}))
    return out


def _continues(prev: Line, ln: Line, bounds) -> bool:
    if ln.bullet or prev.page != ln.page or prev.column != ln.column or len(ln.cells) > 1 or len(prev.cells) > 1:
        return False
    if ln.y0 - prev.y1 > 0.9 * max(ln.size, prev.size):          # a paragraph break, not a wrap
        return False
    if prev.bullet and ln.x0 > prev.x0 + 2:                       # indented under the bullet's text
        return True
    if prev.bullet and not ln.text[:1].islower():                 # not indented: only a lower-case continuation
        return False
    if ln.block != prev.block or (ln.lead_bold and not prev.bold) or _LABEL.match(ln.text):
        return False
    if abs(ln.size - prev.size) > 0.6 or ln.color != prev.color:
        return False
    lo, hi = bounds[(prev.page, prev.column)]
    first_word = ln.text.split(" ", 1)[0]
    # It wrapped: the line ran to the column's edge, or the next word would not have fitted after it.
    return prev.x1 >= hi - 0.15 * (hi - lo) or prev.x1 + 0.55 * ln.size * (len(first_word) + 1) >= hi - 2


_LABEL = re.compile(r"^[A-Z][\w&/+. -]{1,40}:\s")
