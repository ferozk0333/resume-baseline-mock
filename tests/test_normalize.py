"""Document normalization and the extraction runaway guard (iteration 2, phase 9 blocker 1)."""
import io
import re
from types import SimpleNamespace

import docx
import pytest
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import BaseDocTemplate, Frame, FrameBreak, PageTemplate, Paragraph

from core import db
from core.llm import ClaudeLLM, LLMError, RUNAWAY_MESSAGE, RunawayOutput, is_runaway
from extraction import schemas, store
from extraction.onboard import OnboardingInputs, Upload, run_onboarding
from extraction.validate import QuoteIndex
from ingest import documents
from tests.helpers import FakeLLM, src

ROLE_1 = "Institute for Health Metrics and Evaluation"
BULLETS_1 = ["Shipped an end-to-end RAG solution for researchers using AWS Bedrock and OpenSearch.",
             "Optimized embedding dimensions from 1024 to 512, halving vector storage cost."]
ROLE_2 = "Oracle Financial Services Software"
BULLETS_2 = ["Worked on an anomaly detection model on 10M+ payment records using Isolation Forest.",
             "Consumed streaming data from Kafka into ETL pipelines."]


def _raw_pdf(lines: list[tuple[float, float, str]]) -> bytes:
    """PDF drawn with raw text operators (WinAnsi), so tabs and space runs stay real characters as in Word exports."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    ops = ["BT", "/F1 10 Tf"] + [f"1 0 0 1 {x} {y} Tm ({esc(s)}) Tj" for x, y, s in lines] + ["ET"]
    stream = "\n".join(ops).encode("cp1252")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 5 0 R "
        b"/Resources << /Font << /F1 4 0 R >> >> >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def tab_layout_pdf() -> bytes:
    """A resume laid out with tabs and long space runs (right-aligned locations and dates), like a Word export."""
    lines, y = [(50, 740, "Feroz Example"),
                (50, 724, "Seattle, WA |\tLinkedIn:\tlinkedin.com/in/example\t|\tPortfolio:\texample.github.io\t")], 724
    for employer, title, dates, bullets in [(ROLE_1, "AI Engineer", "Oct. 2025 – Jun. 2026", BULLETS_1),
                                            (ROLE_2, "Software Engineer, Machine Learning", "Jul. 2022 – Jun. 2023",
                                             BULLETS_2)]:
        y -= 24
        lines.append((50, y, employer + " " * 60 + "Seattle, WA"))
        y -= 13
        lines.append((50, y, title + "\t\t\t\t\t\t" + " " * 40 + dates))
        for b in bullets:
            y -= 13
            lines.append((58, y, "•\t" + b))
    return _raw_pdf(lines)


def two_column_pdf() -> bytes:
    """Sidebar (skills, education) on the left, experience with bullets on the right."""
    st = getSampleStyleSheet()
    buf = io.BytesIO()
    w, h = LETTER
    frames = [Frame(40, 40, 170, h - 80, id="side"), Frame(230, 40, w - 270, h - 80, id="main")]
    story = [Paragraph("Skills", st["Heading3"]),
             Paragraph("Python, FastAPI, PostgreSQL, Docker, Kubernetes, AWS, REST APIs, CI/CD", st["BodyText"]),
             Paragraph("Education", st["Heading3"]),
             Paragraph("B.S. Computer Engineering, Peachtree Technical University, 2021", st["BodyText"]),
             FrameBreak(), Paragraph("Experience", st["Heading3"])]
    for employer, bullets in [(ROLE_1, BULLETS_1), (ROLE_2, BULLETS_2)]:
        story.append(Paragraph(f"<b>Engineer</b> — {employer} · 2023 – Present", st["BodyText"]))
        story += [Paragraph(b, st["BodyText"], bulletText="•") for b in bullets]
    doc = BaseDocTemplate(buf, pagesize=LETTER)
    doc.addPageTemplates([PageTemplate(id="two", frames=frames)])
    doc.build(story)
    return buf.getvalue()


# --- Normalization ----------------------------------------------------------------------------

def test_tab_layout_is_collapsed_and_keeps_line_structure():
    text = documents.normalize_document("r.pdf", tab_layout_pdf())
    assert "\t" not in text and "  " not in text
    lines = text.splitlines()
    assert any(line.startswith(ROLE_1) and line.endswith("Seattle, WA") for line in lines)
    assert any(line.startswith("AI Engineer") and line.endswith("Oct. 2025 – Jun. 2026") for line in lines)
    for b in BULLETS_1 + BULLETS_2:                    # each bullet on its own line
        assert any(b in line and len(line) < len(b) + 6 for line in lines), b


def _between(text: str, start: str, end: str | None) -> str:
    i = text.index(start)
    return text[i:text.index(end, i)] if end else text[i:]


def test_two_column_pdf_keeps_bullets_with_their_role():
    text = documents.normalize_document("r.pdf", two_column_pdf())
    flat = re.sub(r"\s+", " ", text)
    role_1, role_2 = _between(flat, ROLE_1, ROLE_2), _between(flat, ROLE_2, None)
    assert all(b in role_1 for b in BULLETS_1) and not any(b in role_1 for b in BULLETS_2)
    assert all(b in role_2 for b in BULLETS_2)
    assert "Python, FastAPI" not in role_1 + role_2           # sidebar not interleaved with experience
    assert flat.index("Python, FastAPI") < flat.index(ROLE_1)  # column order: sidebar, then main


def test_markdown_emphasis_links_and_entities_removed():
    md = ("# **Jane**\n\n- Shipped an **end-to-end RAG** solution with _semantic_ chunking and <u>docs</u>\n"
          "continued on the next line &amp; more\n\nAWS cert [, Amazon (Link) 2026](https://www.credly.com/b/1)\n"
          "[me@x.com](mailto:me@x.com) payout_scheduling_service")
    out = documents.clean_markdown(md)
    assert "**" not in out and "<u>" not in out and "&amp;" not in out
    assert "- Shipped an end-to-end RAG solution with semantic chunking and docs continued on the next line & more" \
        in out
    assert "(https://www.credly.com/b/1)" in out              # verify links are kept as text
    assert "me@x.com" in out and "mailto" not in out
    assert "payout_scheduling_service" in out                  # snake_case is not emphasis


def test_wrapped_bullet_lines_join_but_next_header_does_not():
    md = ("- Shipped a RAG solution from ingestion through retrieval to response\ngeneration on AWS.\n"
          "Analytics Intern, Beta Corp, 2022\n- Built REST APIs\nSoftware Engineer I - Harbor, 2021")
    lines = documents.clean_markdown(md).splitlines()
    assert lines == ["- Shipped a RAG solution from ingestion through retrieval to response generation on AWS.",
                     "Analytics Intern, Beta Corp, 2022", "- Built REST APIs", "Software Engineer I - Harbor, 2021"]


def test_docx_headings_and_bullets():
    d = docx.Document()
    d.add_heading("Experience", level=1)
    d.add_paragraph("Data Analyst — Acme · 2023 – Present")
    d.add_paragraph("Built a churn model in Python and scikit-learn that cut monthly churn by 12 percent.",
                    style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    text = documents.normalize_document("r.docx", buf.getvalue())
    assert "## Experience" in text and "- Built a churn model in Python and scikit-learn" in text


def test_normalized_text_is_stored_by_file_hash(tmp_path, monkeypatch):
    conn = db.init_app_db(str(tmp_path / "n.db"))
    data = tab_layout_pdf()
    first = documents.extract_text("r.pdf", data, conn)
    monkeypatch.setattr(documents, "normalize_document", lambda *a: pytest.fail("parsed twice"))
    assert documents.extract_text("again.pdf", data, conn) == first
    row = conn.execute("SELECT normalizer FROM document_texts WHERE file_sha256 = ?",
                       (documents.sha256(data),)).fetchone()
    assert row["normalizer"] == documents.normalizer()


def test_quotes_from_the_normalized_text_pass_the_check():
    text = documents.normalize_document("r.pdf", tab_layout_pdf())
    index = QuoteIndex(text)
    assert all(index.contains(b) for b in BULLETS_1 + BULLETS_2)
    assert index.contains("AI Engineer Oct. 2025 – Jun. 2026")       # header line, spaces collapsed
    assert index.contains("LinkedIn: linkedin.com/in/example")         # tabs became single spaces
    assert not index.contains("Shipped an end-to-end RAG platform")    # paraphrase still fails


def test_tab_layout_resume_extracts_end_to_end(tmp_path):
    """Onboarding stores the normalized text and every quote taken from it survives validation."""
    out = schemas.ProfileExtract(
        roles=[schemas.Role(title="AI Engineer", employer=ROLE_1, industry=None, start="2025-10", end="2026-06",
                            is_current=False, team_size=None,
                            bullets=[schemas.Bullet(text=b, source=src(b)) for b in BULLETS_1],
                            source=src("AI Engineer Oct. 2025 – Jun. 2026")),
               schemas.Role(title="Software Engineer, Machine Learning", employer=ROLE_2, industry=None,
                            start="2022-07", end="2023-06", is_current=False, team_size=None,
                            bullets=[schemas.Bullet(text=b, source=src(b)) for b in BULLETS_2],
                            source=src("Software Engineer, Machine Learning Jul. 2022 – Jun. 2023"))],
        projects=[], claims=[], certifications=[])
    conn = db.init_app_db(str(tmp_path / "e.db"))
    llm = FakeLLM({"resume": out})
    result = run_onboarding(conn, OnboardingInputs(resume=Upload("r.pdf", tab_layout_pdf())), llm)
    sent = llm.calls[0]["user"]
    assert "\t" not in sent and "  " not in sent                       # the model sees the normalized text
    profile = store.load_profile(conn, result.candidate_id)
    assert [len(r.bullets) for r in profile.roles] == [2, 2]
    assert store.list_dropped(conn, result.candidate_id) == []
    stored = conn.execute("SELECT raw_text FROM evidence_items WHERE type = 'resume'").fetchone()[0]
    assert stored in sent                                               # same text for extraction and quotes


# --- Runaway guard ---------------------------------------------------------------------------

def test_is_runaway():
    assert is_runaway('{"quote": "AI Engineer' + "\\t" * 100)
    assert is_runaway("x" + " " * 300)
    assert not is_runaway('{"roles": [{"title": "Engineer", "bullets": ["a", "b"]}], "claims": []}' * 3)
    assert not is_runaway("\\t" * 20)                                    # short runs are fine


class StreamingMessages:
    """client.messages with `stream`: each call plays the next scripted text, in small deltas."""

    def __init__(self, scripts: list[str]):
        self.scripts = scripts
        self.calls = 0
        self.closed = 0

    def stream(self, **request):
        text = self.scripts[self.calls]
        self.calls += 1
        outer = self

        class Stream:
            def __enter__(self):
                return self

            def __exit__(self, *exc):
                outer.closed += 1

            def __iter__(self):
                for i in range(0, len(text), 8):
                    yield SimpleNamespace(type="content_block_delta",
                                          delta=SimpleNamespace(type="text_delta", text=text[i:i + 8]))

            def get_final_message(self):
                parsed = request["output_format"].model_validate_json(text)
                return SimpleNamespace(stop_reason="end_turn", parsed_output=parsed,
                                       usage=SimpleNamespace(input_tokens=10, output_tokens=20))
        return Stream()


GOOD = '{"projects": [], "claims": []}'
RUNAWAY = '{"projects": [{"name": "' + "\\t" * 400


def test_runaway_output_is_retried_once():
    msgs = StreamingMessages([RUNAWAY, GOOD])
    out = ClaudeLLM(client=SimpleNamespace(messages=msgs)).parse(
        model="claude-haiku-4-5", system="S", user="U", output_type=schemas.PortfolioExtract)
    assert out.projects == [] and msgs.calls == 2 and msgs.closed == 2


def test_runaway_twice_fails_with_a_clear_message():
    msgs = StreamingMessages([RUNAWAY, RUNAWAY])
    with pytest.raises(RunawayOutput) as exc:
        ClaudeLLM(client=SimpleNamespace(messages=msgs)).parse(
            model="claude-haiku-4-5", system="S", user="U", output_type=schemas.PortfolioExtract)
    assert str(exc.value) == RUNAWAY_MESSAGE and isinstance(exc.value, LLMError)


def test_runaway_twice_marks_the_source_failed_with_a_warning(tmp_path):
    class Runaway(FakeLLM):
        def parse(self, *, model, system, user, output_type):
            raise RunawayOutput(RUNAWAY_MESSAGE)

    conn = db.init_app_db(str(tmp_path / "r.db"))
    result = run_onboarding(conn, OnboardingInputs(resume=Upload("r.pdf", tab_layout_pdf())), Runaway({}))
    row = conn.execute("SELECT status, error FROM evidence_items WHERE type = 'resume'").fetchone()
    assert row["status"] == "failed" and RUNAWAY_MESSAGE in row["error"]
    assert any(RUNAWAY_MESSAGE in w for w in result.warnings)


def test_identical_document_is_extracted_once_but_a_rerun_asks_again(tmp_path, monkeypatch):
    from core import config
    from graph.context import Deps
    from graph.pipeline import Pipeline
    from tests.helpers import SAMPLE_RESUME_LINES, make_pdf, sample_resume_extract
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    path = str(tmp_path / "x.db")
    db.init_app_db(path).close()
    llm = FakeLLM({"resume": sample_resume_extract()})
    pipe = Pipeline(Deps(path, judge=None, llm=llm, extraction_cache=True))
    upload = Upload("cv.pdf", make_pdf(SAMPLE_RESUME_LINES))
    first = pipe.start_onboarding(OnboardingInputs(resume=upload), "baseline-data-scientist-entry", early=False)
    pipe.start_onboarding(OnboardingInputs(resume=upload), "baseline-data-scientist-entry", early=False)
    assert len(llm.calls) == 1                                    # second upload: from the cache
    pipe.start_onboarding(OnboardingInputs(resume=upload, candidate_id=first["candidate_id"]),
                          "baseline-data-scientist-entry", early=False)
    assert len(llm.calls) == 2                                    # "Re-run extraction": fresh call
